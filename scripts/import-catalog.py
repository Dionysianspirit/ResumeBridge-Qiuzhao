"""Read a community catalog workbook without modifying it; retain every source row.

Personal matching columns (挂钩提示, 软岗分) are cleared in the JSON output so a
local annotated 土豆-style workbook can be published or imported without those notes.
"""
import argparse
import hashlib
import json
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from openpyxl import load_workbook

parser = argparse.ArgumentParser()
parser.add_argument("source", type=Path)
parser.add_argument("destination", type=Path)
args = parser.parse_args()
if args.destination.exists():
    raise SystemExit("Destination already exists; refusing to overwrite")
source_hash = hashlib.sha256(args.source.read_bytes()).hexdigest()
book = load_workbook(args.source, read_only=True, data_only=False)

def value(v):
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    return "" if v is None else v

def read_rows(name):
    sheet = book[name]
    headers = [c.value for c in sheet[1]]
    result = []
    for row in sheet.iter_rows(min_row=2):
        if not any(c.value is not None for c in row):
            continue
        if any(c.data_type == "f" for c in row):
            raise ValueError(f"Unexpected formula in {name}!{row[0].row}")
        result.append((row[0].row, dict(zip(headers, [value(c.value) for c in row]))))
    return result

def key(fields):
    return tuple(str(fields.get(k, "")).strip() for k in ["公司名称", "批次", "投递方式", "招聘岗位"])

weekly = read_rows("本周加急约100")
weekly_by_key = {key(fields): (row, fields) for row, fields in weekly}
entries = []
lookup = {}
row_counts = {}
for sheet_name, source_expired in [("总筛未过期", False), ("已过期剔除", True)]:
    source_rows = read_rows(sheet_name)
    row_counts[sheet_name] = len(source_rows)
    for row, fields in source_rows:
        identity = (source_expired, *key(fields))
        source = {"sheet": sheet_name, "row": row}
        if identity in lookup:
            entry = lookup[identity]
            entry["sources"].append(source)
            entry["variants"].append({**source, "fields": fields})
            continue
        digest = hashlib.sha256(json.dumps(identity, ensure_ascii=False).encode("utf-8")).hexdigest()[:24]
        weekly_source = weekly_by_key.get(key(fields)) if not source_expired else None
        entry = {
            "id": "catalog-" + digest,
            "companyName": fields["公司名称"],
            "roles": fields["招聘岗位"],
            "location": fields["工作地点"],
            "batch": fields["批次"],
            "applyUrl": fields["投递方式"],
            "noticeUrl": fields["官方公告"],
            "tier": fields["档位"],
            "tierLabel": fields["档位含义"],
            "deadline": fields["截止日期"],
            "deadlineType": fields["截止类型"],
            "sourceExpired": source_expired,
            "weeklyRank": int(weekly_source[1]["本周序"]) if weekly_source else None,
            "weeklyFields": weekly_source[1] if weekly_source else None,
            "fields": fields,
            "sources": [source],
            "variants": [],
        }
        entries.append(entry)
        lookup[identity] = entry

assert len(weekly_by_key) == len(weekly) == 100
assert sum(e["weeklyRank"] is not None for e in entries) == len(weekly)
assert sum(len(e["sources"]) for e in entries) == sum(row_counts.values())
assert len({e["id"] for e in entries}) == len(entries)
counts = {
    "activeSourceRows": row_counts["总筛未过期"],
    "expiredSourceRows": row_counts["已过期剔除"],
    "active": sum(not e["sourceExpired"] for e in entries),
    "expired": sum(e["sourceExpired"] for e in entries),
    "weekly": len(weekly),
    "mergedSourceRows": sum(len(e["sources"]) - 1 for e in entries),
}
PERSONAL_KEYS = ("挂钩提示", "软岗分")

def strip_personal(value):
    if isinstance(value, dict):
        return {key: "" if key in PERSONAL_KEYS else strip_personal(item) for key, item in value.items()}
    if isinstance(value, list):
        return [strip_personal(item) for item in value]
    return value

safe_name = args.source.name
if any(token in safe_name for token in ("李-", "澪", "ASUS")):
    safe_name = "imported-catalog.json"
result = strip_personal({
    "schemaVersion": 1,
    "sourceFile": safe_name,
    "sourceSha256": source_hash,
    "asOf": "2026-09-06",
    "counts": counts,
    "entries": entries,
})
args.destination.parent.mkdir(parents=True, exist_ok=True)
args.destination.write_text(json.dumps(result, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
book.close()
assert hashlib.sha256(args.source.read_bytes()).hexdigest() == source_hash
print(json.dumps({"counts": counts, "bytes": args.destination.stat().st_size, "sourceUnchanged": True}, ensure_ascii=False))
